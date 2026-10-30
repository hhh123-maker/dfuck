import sqlite3
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from datetime import datetime

conn = sqlite3.connect("attendance.db")
cur = conn.cursor()
cur.execute("SELECT id, student_name, sign_time FROM sign_record ORDER BY sign_time DESC")
rows = cur.fetchall()
conn.close()

wb = Workbook()
ws = wb.active
ws.title = "签到记录"

headers = ["序号", "学生姓名", "签到时间"]
ws.append(headers)

header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
header_font = Font(bold=True, color="FFFFFF", size=12)
for col in range(1, len(headers)+1):
    cell = ws.cell(row=1, column=col)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center")

seq = 1
for r in rows:
    stu_id = r[1]
    t = r[2]
    ws.append([seq, name, t])
    seq += 1

ws.column_dimensions['A'].width = 8
ws.column_dimensions['B'].width = 15
ws.column_dimensions['C'].width = 25

now_str = datetime.now().strftime("%Y%m%d_%H%M%S")
filename = f"签到记录_{now_str}.xlsx"

wb.save(filename)
print(f"✅ 已导出 {len(rows)} 条签到记录 → {filename}")